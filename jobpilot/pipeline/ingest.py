"""Stage 0: resume file -> MasterResume (data/master_resume.yaml).

Run once, then *read and edit the YAML by hand*. It is the only thing the
tailoring step is allowed to draw facts from, so this is where you add the
bullets that didn't fit on your one-page resume: the more evidence here, the
better every tailored version gets.
"""

from __future__ import annotations

from pathlib import Path

from ..llm import StructuredLLM
from ..models import MasterResume

SYSTEM = """You convert a resume into structured data.

Rules:
- Transcribe faithfully. Keep each bullet's wording; fix only obvious typos.
- Never invent, infer or embellish facts, numbers, dates, titles or skills.
- For each bullet, list in `skills` only the tools/skills that bullet explicitly names
  or unambiguously demonstrates (e.g. "built a React dashboard" -> React).
- Leave `id` fields empty; they are assigned afterwards.
- Dates: keep the resume's precision (e.g. "2021-06", "Jun 2021", "2021"). Use "Present" for current roles."""


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    return path.read_text()


def ingest_resume(path: Path, llm: StructuredLLM) -> MasterResume:
    text = extract_text(path)
    if len(text.strip()) < 200:
        raise ValueError(
            f"Only {len(text.strip())} characters of text found in {path}. "
            "If it's a scanned/image PDF, export it as text or DOCX first."
        )
    resume = llm.parse(system=SYSTEM, prompt=f"<resume>\n{text}\n</resume>", schema=MasterResume, effort="medium")
    return resume.assign_ids()
